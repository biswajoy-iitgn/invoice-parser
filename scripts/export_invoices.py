#!/usr/bin/env python3
"""One command: every invoice in the database -> an Excel workbook. No row limit.

  macOS / Linux :  ./export_invoices.sh                 (or: python scripts/export_invoices.py)
  Windows       :  export_invoices.bat                  (double-click works too)

  Options:  --from 2026-04-01  --to 2026-06-30  --vendor "dachser"  --invoice 7044012289
            --out path.xlsx    --driver auto|odbc|native    --dsn NAME    --no-audit

It reads the reporting view `reporting.invoice_export` with a free PostgreSQL driver:
  * odbc   : psqlODBC through pyodbc ("psqlODBCw" on macOS/Linux, "PostgreSQL Unicode(x64)" on Windows)
  * native : psycopg (libpq), used when no psqlODBC driver is installed
`--driver auto` (default) tries ODBC first, then native. Neither has the 3-row cap of an unlicensed
Excel driver. The database address comes from DATABASE_URL_READONLY or DATABASE_URL, taken from the
environment or from `.env.local` next to the project (never printed, never written to the workbook).

Deliberately standalone (standard library + openpyxl + one driver) so it also runs on a Windows PC
that has none of the extraction stack installed.
"""
from __future__ import annotations

import argparse
import datetime as dt
import decimal
import getpass
import hashlib
import json
import os
import platform
import re
import sys
import uuid
from pathlib import Path
from urllib.parse import unquote, urlparse

ROOT = Path(__file__).resolve().parents[1]
VIEW = "reporting.invoice_export"

# view column -> the heading used everywhere else in this project (same order as the JSON output)
HEADERS = [
    ("invoice_number", "Invoice No"), ("invoice_date", "Invoice Date"), ("container_number", "Container Number"),
    ("mawb_no", "MAWB No"), ("hawb_no", "HAWB No"), ("vendor_name", "Vendor Name"),
    ("bill_of_lading_hbl", "Bill of Lading / HBL"), ("weight", "WEIGHT"),
    ("sea_freight", "Sea Freight"), ("air_freight", "Air Freight"),
    ("terminal_handling_charges", "Terminal Handling Charges"), ("airway_waybill_charges", "Air Waybill Charges"),
    ("bl_charges", "BL Charges"), ("ens_charges", "ENS Charges"), ("seal_charges", "Seal Charges"), ("vgm", "VGM"),
    ("edi_fees", "EDI Fees"), ("dap_charges", "DAP Charges"), ("other_charges", "Other Charges"),
    ("ams_charges", "AMS CHARGES"), ("handling_charges", "Handling Charges"),
    ("pre_carriage_haulage", "Pre Carriage Haulage"), ("ex_rate", "Ex Rate"), ("misc", "Misc"),
    ("transportation", "Transportation"), ("warehouse_charges_port", "Warehouse charges Port"),
    ("examination_charges", "Examination charges"), ("fuel_charges", "Fuel Charges"),
    ("cargo_screening", "Cargo Screening"), ("agency_charges", "Agency Charges"),
    ("shipping_bill_charges", "Shipping bill charges"),
]
STATUS = [("review_status", "Review status"), ("extraction_status", "Extraction status"),
          ("updated_at", "Last updated")]
CHARGES = [c for c, _ in HEADERS[8:] if c != "ex_rate"]
TITLE = dict(HEADERS + STATUS)


def load_env_file() -> None:
    """KEY='value' lines from .env.local (same file scripts/run_app.sh reads); the environment wins."""
    f = ROOT / ".env.local"
    if not f.exists():
        return
    for line in f.read_text().splitlines():
        m = re.match(r"\s*([A-Z_][A-Z0-9_]*)\s*=\s*(.*)\s*$", line)
        if m and m.group(1) not in os.environ:
            os.environ[m.group(1)] = m.group(2).strip().strip("'\"")


def where(a) -> tuple[str, list]:
    cond, args = [], []
    if a.date_from:
        cond.append("invoice_date >= ?")
        args.append(dt.date.fromisoformat(a.date_from))
    if a.date_to:
        cond.append("invoice_date <= ?")
        args.append(dt.date.fromisoformat(a.date_to))
    if a.vendor:
        cond.append("vendor_name ILIKE ?")
        args.append(f"%{a.vendor}%")
    if a.invoice:
        cond.append("invoice_number = ?")
        args.append(a.invoice)
    return (" WHERE " + " AND ".join(cond)) if cond else "", args


def fetch_odbc(url: str | None, dsn: str | None, sql: str, args: list):
    import pyodbc

    if dsn:
        cs = f"DSN={dsn}"
        if os.environ.get("ODBC_UID"):
            cs += f";UID={os.environ['ODBC_UID']};PWD={os.environ.get('ODBC_PWD', '')}"
        name = f"ODBC DSN {dsn}"
    else:
        drivers = [d for d in pyodbc.drivers() if re.search(r"postgres|psqlodbc", d, re.I)]
        if not drivers:
            raise RuntimeError("no psqlODBC driver installed")
        drv = next((d for d in drivers if re.search(r"unicode|odbcw", d, re.I)), drivers[0])
        u = urlparse(url)
        cs = (f"Driver={{{drv}}};Server={u.hostname};Port={u.port or 5432};Database={u.path.lstrip('/')};"
              f"UID={unquote(u.username or '')};PWD={unquote(u.password or '')};SSLmode=require")
        name = f"psqlODBC ({drv}) via pyodbc"
    cn = pyodbc.connect(cs, autocommit=True, timeout=30)
    cur = cn.cursor()
    cur.execute(sql, args)
    cols = [c[0] for c in cur.description]
    rows = [dict(zip(cols, r)) for r in cur.fetchall()]
    cn.close()
    return rows, cols, name


def fetch_native(url: str, sql: str, args: list):
    import psycopg

    if "sslmode=" not in url and (urlparse(url).hostname or "") not in ("localhost", "127.0.0.1", "::1", ""):
        url += ("&" if "?" in url else "?") + "sslmode=require"
    with psycopg.connect(url, connect_timeout=30) as cn, cn.cursor() as cur:
        cur.execute(sql.replace("?", "%s"), args)
        cols = [c.name for c in cur.description]
        rows = [dict(zip(cols, r)) for r in cur.fetchall()]
    return rows, cols, f"psycopg {psycopg.__version__} (libpq)"


def cell(v):
    if isinstance(v, decimal.Decimal):
        return float(v)
    if isinstance(v, dt.datetime):
        return v.replace(tzinfo=None)
    return v


def build_workbook(rows: list[dict], info: dict, out: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    wb = Workbook()

    def sheet(ws, columns: list[tuple[str, str]], data: list[dict], table: str, money: set[str], dates: dict[str, str]):
        ws.append([t for _, t in columns])
        for r in data:
            ws.append([cell(r.get(c)) for c, _ in columns])
        for i, (c, t) in enumerate(columns, 1):
            col = get_column_letter(i)
            h = ws.cell(row=1, column=i)
            h.font, h.fill, h.alignment = head_font, head_fill, Alignment(wrap_text=True, vertical="center")
            width = max([len(t)] + [len(str(r.get(c) if r.get(c) is not None else "")) for r in data[:400]])
            ws.column_dimensions[col].width = min(44, max(11, width + 2))
            fmt = "#,##0.00" if c in money else dates.get(c)
            if fmt:
                for row in range(2, len(data) + 2):
                    ws.cell(row=row, column=i).number_format = fmt
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 32
        if data:
            t = Table(displayName=table, ref=f"A1:{get_column_letter(len(columns))}{len(data) + 1}")
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)

    ws = wb.active
    ws.title = "Invoices"
    sheet(ws, HEADERS + STATUS, rows, "Invoices", set(CHARGES),
          {"invoice_date": "dd-mmm-yyyy", "updated_at": "dd-mmm-yyyy hh:mm"})

    long = [{"invoice_number": r.get("invoice_number"), "invoice_date": r.get("invoice_date"),
             "vendor_name": r.get("vendor_name"), "category": TITLE[c], "amount": r.get(c)}
            for r in rows for c in CHARGES if r.get(c) not in (None, "")]
    sheet(wb.create_sheet("Charges"), [("invoice_number", "Invoice No"), ("invoice_date", "Invoice Date"),
                                       ("vendor_name", "Vendor Name"), ("category", "Charge category"),
                                       ("amount", "Amount (INR)")], long, "Charges", {"amount"},
          {"invoice_date": "dd-mmm-yyyy"})

    def total(r):
        return sum(float(r[c]) for c in CHARGES if r.get(c) not in (None, ""))

    vend: dict[str, list] = {}
    for r in rows:
        v = vend.setdefault(r.get("vendor_name") or "(no vendor)", [0, 0.0])
        v[0] += 1
        v[1] += total(r)
    by_vendor = [{"vendor": k, "invoices": v[0], "total": round(v[1], 2)}
                 for k, v in sorted(vend.items(), key=lambda kv: -kv[1][1])]
    sheet(wb.create_sheet("By vendor"), [("vendor", "Vendor Name"), ("invoices", "Invoices"),
                                         ("total", "Total charges (INR)")], by_vendor, "ByVendor", {"total"}, {})
    by_cat = [{"category": TITLE[c], "invoices": sum(1 for r in rows if r.get(c) not in (None, "")),
               "total": round(sum(float(r[c]) for r in rows if r.get(c) not in (None, "")), 2)} for c in CHARGES]
    by_cat = sorted([b for b in by_cat if b["invoices"]], key=lambda b: -b["total"])
    sheet(wb.create_sheet("By category"), [("category", "Charge category"), ("invoices", "Invoices with it"),
                                           ("total", "Total (INR)")], by_cat, "ByCategory", {"total"}, {})

    ws = wb.create_sheet("Export info")
    for k, v in info.items():
        ws.append([k, v])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 90
    for row in range(1, len(info) + 1):
        ws.cell(row=row, column=1).font = Font(bold=True)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main() -> int:
    ap = argparse.ArgumentParser(description="Export all invoices from the database to Excel (no row limit).")
    ap.add_argument("--from", dest="date_from", help="invoice date from (YYYY-MM-DD)")
    ap.add_argument("--to", dest="date_to", help="invoice date to (YYYY-MM-DD)")
    ap.add_argument("--vendor", help="vendor name contains")
    ap.add_argument("--invoice", help="one invoice number")
    ap.add_argument("--out", help="output .xlsx (default exports/invoices_<timestamp>.xlsx)")
    ap.add_argument("--driver", choices=["auto", "odbc", "native"], default="auto")
    ap.add_argument("--dsn", help="use this ODBC data source instead of the database URL")
    ap.add_argument("--no-audit", action="store_true", help="do not record this export in the audit log")
    a = ap.parse_args()
    load_env_file()
    url = os.environ.get("DATABASE_URL_READONLY") or os.environ.get("DATABASE_URL")
    if not url and not a.dsn:
        print("No database address: set DATABASE_URL_READONLY (or DATABASE_URL) in the environment or in .env.local,"
              " or pass --dsn NAME.", file=sys.stderr)
        return 2
    w, args = where(a)
    sql = f"SELECT * FROM {VIEW}{w} ORDER BY invoice_date NULLS LAST, invoice_number"
    rows = cols = used = None
    errors = []
    for drv in (["odbc", "native"] if a.driver == "auto" else [a.driver]):
        try:
            if drv == "odbc":
                rows, cols, used = fetch_odbc(url, a.dsn, sql, args)
            else:
                rows, cols, used = fetch_native(url, sql, args)
            break
        except Exception as exc:      # try the next driver; report all if none works
            errors.append(f"{drv}: {type(exc).__name__}: {str(exc)[:160]}")
    if rows is None:
        print("Could not read the database:\n  " + "\n  ".join(errors), file=sys.stderr)
        return 1
    missing = [c for c, _ in HEADERS if c not in cols]
    if missing:
        print(f"The view is missing expected columns: {missing}", file=sys.stderr)
        return 1
    now = dt.datetime.now()
    out = Path(a.out) if a.out else ROOT / "exports" / f"invoices_{now:%Y%m%d_%H%M%S}.xlsx"
    digest = hashlib.sha256(json.dumps(rows, sort_keys=True, default=str).encode()).hexdigest()
    host = urlparse(url).hostname if url else f"DSN {a.dsn}"
    filters = {k: v for k, v in (("from", a.date_from), ("to", a.date_to), ("vendor", a.vendor),
                                 ("invoice", a.invoice)) if v}
    info = {"Exported at": now.strftime("%Y-%m-%d %H:%M:%S"), "Exported by": getpass.getuser(),
            "Computer": platform.node(), "Database host": host, "Source": VIEW, "Driver": used,
            "Filters": json.dumps(filters) if filters else "none (all invoices)", "Invoices": len(rows),
            "Charge rows": sum(1 for r in rows for c in CHARGES if r.get(c) not in (None, "")),
            "Data fingerprint (SHA-256)": digest,
            "Note": "Blank cells mean the field is not on the invoice. Amounts are INR before GST."}
    build_workbook(rows, info, out)
    audited = "not recorded"
    if not a.no_audit and os.environ.get("DATABASE_URL") and os.environ.get("INVOICE_ORG_ID"):
        try:      # needs a role that may write audit_logs; the read-only reporting role cannot, which is fine
            import psycopg

            with psycopg.connect(os.environ["DATABASE_URL"], connect_timeout=30) as cn:
                cn.execute("SELECT set_config('app.current_org', %s, TRUE)", (os.environ["INVOICE_ORG_ID"],))
                cn.execute("INSERT INTO audit_logs (organization_id, entity, entity_id, action, new_json) "
                           "VALUES (%s, 'export', %s, 'EXPORT', %s)",
                           (os.environ["INVOICE_ORG_ID"], str(uuid.uuid4()),
                            json.dumps({"file": out.name, "rows": len(rows), "sha256": digest, "driver": used,
                                        "filters": filters, "by": getpass.getuser(), "host": platform.node()})))
            audited = "recorded in audit_logs"
        except Exception as exc:
            audited = f"not recorded ({type(exc).__name__})"
    print(f"{len(rows)} invoices -> {out}\n  driver: {used}\n  audit: {audited}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
