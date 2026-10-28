"""MIS workbook: the analytics snapshot as an Excel file with native charts (for finance / management)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FONT, HEAD_FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
MONEY = "#,##0.00"
SEVERITY_FILL = {"high": "F8D7DA", "medium": "FDEBD0", "low": "EEF1F5"}


def _table(ws, headers: list[str], rows: list[list], money: tuple[int, ...] = (), pct: tuple[int, ...] = (),
           start_row: int = 1, widths: dict[int, int] | None = None) -> int:
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font, c.fill, c.alignment = HEAD_FONT, HEAD_FILL, Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, start_row + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            if j in money:
                c.number_format = MONEY
            elif j in pct:
                c.number_format = '0.0"%"'
    for j, h in enumerate(headers, 1):
        longest = max([len(str(h))] + [len(str(r[j - 1])) for r in rows[:300] if j - 1 < len(r) and r[j - 1] is not None])
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, min(60, max(12, longest + 2)))
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return start_row + len(rows)


def build(snap: dict, out: Path) -> Path:
    wb = Workbook()
    m = snap["meta"]

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Invoice MIS report"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = (f"Period {m['period_from'] or '—'} to {m['period_to'] or '—'} · generated {m['generated_at']} · "
                f"{m['currency']}")
    ws["A2"].font = Font(color="666666")
    end = _table(ws, ["Measure", "Value", "Note"],
                 [[k["label"], k["value"], k.get("note", "")] for k in snap["kpis"]], start_row=4, widths={1: 34, 2: 22, 3: 70})
    for r in range(5, end + 1):
        if isinstance(ws.cell(row=r, column=2).value, float):
            ws.cell(row=r, column=2).number_format = MONEY
    ws.cell(row=end + 2, column=1, value="What the numbers say").font = Font(bold=True, size=12)
    for i, line in enumerate(snap["narrative"], end + 3):
        ws.cell(row=i, column=1, value=f"• {line}").alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
        ws.row_dimensions[i].height = 30
    ws.freeze_panes = None

    ws = wb.create_sheet("Monthly")
    rows = [[x["month"], x["invoices"], x["spend"], "actual"] for x in snap["monthly"]]
    if snap["forecast"].get("available"):
        rows += [[p["month"], None, p["spend"], f"forecast ({p['low']:,.0f} – {p['high']:,.0f})"] for p in snap["forecast"]["points"]]
    end = _table(ws, ["Month", "Invoices", "Charges before GST (INR)", "Kind"], rows, money=(3,))
    if len(snap["monthly"]) >= 2:
        ch = LineChart()
        ch.title, ch.height, ch.width, ch.legend = "Monthly charges (INR)", 8, 22, None
        n = len(snap["monthly"])
        ch.add_data(Reference(ws, min_col=3, min_row=1, max_row=n + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
        ws.add_chart(ch, "F2")

    ws = wb.create_sheet("Vendors")
    rows = [[v["vendor"], v["invoices"], v["spend"], v["share"], v["avg_invoice"]] for v in snap["vendors"]]
    _table(ws, ["Vendor", "Invoices", "Charges (INR)", "Share of spend", "Average invoice (INR)"], rows, money=(3, 5), pct=(4,))
    if rows:
        ch = BarChart()
        ch.type, ch.title, ch.height, ch.width, ch.legend = "bar", "Largest vendors by charges (INR)", 9, 22, None
        k = min(10, len(rows))
        ch.add_data(Reference(ws, min_col=3, min_row=1, max_row=k + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=k + 1))
        ch.y_axis.scaling.orientation = "minMax"
        ch.x_axis.scaling.orientation = "maxMin"        # largest on top
        ws.add_chart(ch, "H2")

    ws = wb.create_sheet("Categories")
    rows = [[c["category"], c["invoices"], c["spend"], c["share"], c["avg"]] for c in snap["categories"]]
    _table(ws, ["Charge category", "Invoices with it", "Total (INR)", "Share of spend", "Average when charged (INR)"],
           rows, money=(3, 5), pct=(4,))
    if rows:
        ch = BarChart()
        ch.type, ch.title, ch.height, ch.width, ch.legend = "bar", "Charges by category (INR)", 10, 22, None
        ch.add_data(Reference(ws, min_col=3, min_row=1, max_row=len(rows) + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(rows) + 1))
        ch.x_axis.scaling.orientation = "maxMin"
        ws.add_chart(ch, "H2")

    ws = wb.create_sheet("Vendor x category")
    vc = snap["vendor_category"]
    _table(ws, ["Vendor"] + vc["categories"] + ["Total"],
           [[r["vendor"]] + r["values"] + [round(sum(r["values"]), 2)] for r in vc["rows"]],
           money=tuple(range(2, len(vc["categories"]) + 3)))

    ws = wb.create_sheet("Freight per kg")
    _table(ws, ["Vendor", "Mode", "Invoices", "Median INR per kg", "Lowest", "Highest"],
           [[u["vendor"], u["mode"], u["invoices"], u["median_per_kg"], u["min_per_kg"], u["max_per_kg"]]
            for u in snap["unit_costs"]], money=(4, 5, 6))

    ws = wb.create_sheet("Findings")
    rows = [[f["severity"].upper(), f["type"].replace("_", " "), f["title"], f["detail"], ", ".join(f["invoices"])]
            for f in snap["findings"]]
    end = _table(ws, ["Priority", "Kind", "Finding", "Detail", "Invoices"], rows, widths={1: 10, 2: 22, 3: 58, 4: 80, 5: 40})
    for r, f in enumerate(snap["findings"], 2):
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=SEVERITY_FILL[f["severity"]])
        for c in (3, 4):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Review queue")
    _table(ws, ["Invoice", "Vendor", "Invoice date", "Charges (INR)"],
           [[r["invoice"], r["vendor"], r["date"], r["total"]] for r in snap["review_queue"]], money=(4,))

    ws = wb.create_sheet("Data quality")
    q = snap.get("quality") or {}
    end = _table(ws, ["Field", "Filled", "Out of", "Fill rate"],
                 [[r["field"], r["filled"], r["of"], r["rate"]] for r in q.get("fill_rates", [])], pct=(4,))
    row = end + 2
    for title, d in (("Review status", q.get("review_status")), ("Extraction status", q.get("extraction_status")),
                     ("Routing of the latest run", (q.get("runs") or {}).get("routes")),
                     ("Pipeline versions", (q.get("runs") or {}).get("pipeline_versions"))):
        if d:
            ws.cell(row=row, column=1, value=title).font = Font(bold=True)
            for k, v in d.items():
                row += 1
                ws.cell(row=row, column=1, value=str(k))
                ws.cell(row=row, column=2, value=v)
            row += 2
    ws.freeze_panes = None

    ws = wb.create_sheet("About")
    info = [("Engine", m["engine"]), ("Generated at", m["generated_at"]), ("Invoices analysed", m["invoices"]),
            ("Period", f"{m['period_from']} to {m['period_to']}"), ("Database host", (m.get("source") or {}).get("host")),
            ("Source view", (m.get("source") or {}).get("view")), ("Input fingerprint (SHA-256)", m["fingerprint"]),
            ("Amounts", m["currency"])]
    if snap["forecast"].get("available"):
        info.append(("Forecast", snap["forecast"]["method"] + (f"; back-test error {snap['forecast']['backtest_error_pct']}%"
                                                              if "backtest_error_pct" in snap["forecast"] else "")))
    else:
        info.append(("Forecast", snap["forecast"].get("reason", "")))
    info += [(f"Method: {k.replace('_', ' ')}", v) for k, v in snap["methods"].items()]
    for i, (k, v) in enumerate(info, 1):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 120
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
