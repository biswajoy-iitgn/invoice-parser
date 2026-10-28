"""Analytics engine: aggregates, findings and reports on small hand-built data (no database needed)."""

import json
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from analytics import insights, report_html, report_xlsx  # noqa: E402
from analytics.data import fingerprint, prepare, weight_kg  # noqa: E402


def frame(rows):
    base = {"invoice_number": None, "invoice_date": None, "vendor_name": None, "container_number": None,
            "mawb_no": None, "hawb_no": None, "bill_of_lading_hbl": None, "weight": None, "ex_rate": None,
            "review_status": "APPROVED", "extraction_status": "COMPLETED"}
    return prepare(pd.DataFrame([{**base, **r} for r in rows]))


def many(n=60):
    rows = []
    for i in range(n):
        month = 1 + i % 10
        rows.append({"invoice_number": f"INV-{i}", "invoice_date": f"2026-{month:02d}-{1 + i % 27:02d}",
                     "vendor_name": ["Acme Freight Pvt Ltd", "Blue Line Logistics", "Coastal Cargo LLP"][i % 3],
                     "container_number": "MSCU6639295", "bill_of_lading_hbl": f"BL{i}", "weight": "1,000 KG",
                     "sea_freight": 100000 + 500 * (i % 7), "terminal_handling_charges": 9000 + 10 * (i % 5),
                     "agency_charges": 1000 + (i % 4), "ex_rate": "USD 84.10",
                     "review_status": "PENDING" if i % 10 == 0 else "APPROVED"})
    return rows


class Aggregates(unittest.TestCase):
    def test_prepare_types_and_mode(self):
        inv = frame([{"invoice_number": "A", "invoice_date": "2026-05-03", "vendor_name": "V", "sea_freight": "100.50",
                      "container_number": "MSCU6639295, TGBU6174371", "weight": "2,905 KG"},
                     {"invoice_number": "B", "invoice_date": "2026-05-09", "vendor_name": "V", "air_freight": 50,
                      "mawb_no": "125-12345675", "weight": "1.5 MT"}])
        self.assertEqual(list(inv["mode"]), ["sea", "air"])
        self.assertEqual(inv["total"].tolist(), [100.5, 50.0])
        self.assertEqual(inv["containers"].tolist(), [2, 0])
        self.assertEqual(inv["weight_kg"].tolist(), [2905.0, 1500.0])
        self.assertEqual(weight_kg("6750.500"), 6750.5)
        self.assertEqual(len(fingerprint(inv)), 64)

    def test_small_data_reports_not_enough(self):
        s = insights.analyse({"invoices": frame(many(5)), "runs": None}, now=pd.Timestamp("2026-12-15"))
        self.assertFalse(s["forecast"]["available"])
        self.assertIn("Needs 6 complete months", s["forecast"]["reason"])
        self.assertTrue(any("fewer than 40 invoices" in line for line in s["narrative"]))
        json.dumps(s)      # the snapshot must be plain JSON

    def test_totals_and_shares(self):
        s = insights.analyse({"invoices": frame(many(60)), "runs": None}, now=pd.Timestamp("2026-12-15"))
        self.assertEqual(s["meta"]["invoices"], 60)
        self.assertAlmostEqual(sum(v["spend"] for v in s["vendors"]), s["kpis"][1]["value"], places=1)
        self.assertAlmostEqual(sum(c["share"] for c in s["categories"]), 100.0, delta=0.3)
        self.assertEqual(s["categories"][0]["category"], "Sea Freight")
        self.assertEqual(len(s["review_queue"]), 6)
        self.assertTrue(s["forecast"]["available"])
        self.assertEqual(len(s["forecast"]["points"]), 3)


class Findings(unittest.TestCase):
    def test_outlier_charge_is_flagged_with_its_invoice(self):
        rows = many(60)
        rows[7]["agency_charges"] = 95000          # everyone else pays about 1,000
        out = insights.outlier_charges(frame(rows))
        hit = [f for f in out if f["invoices"] == ["INV-7"]]
        self.assertEqual(len(hit), 1)
        self.assertEqual(hit[0]["metric"]["category"], "Agency Charges")
        self.assertEqual(hit[0]["severity"], "high")
        self.assertEqual(insights.outlier_charges(frame(many(60))), [])      # nothing flagged in clean data

    def test_duplicates(self):
        rows = many(12)
        rows.append({**rows[3], "vendor_name": "Acme Freight Private Limited"})           # same number, other spelling
        rows.append({**rows[5], "invoice_number": "INV-5-COPY"})                          # same vendor/date/amount
        kinds = {f["type"] for f in insights.duplicates(frame(rows))}
        self.assertIn("duplicate_number", kinds)
        self.assertIn("possible_double_billing", kinds)

    def test_fx_rate_out_of_line(self):
        rows = many(30)
        rows[4]["ex_rate"] = "USD 92.50"
        out = insights.fx_checks(frame(rows))
        self.assertEqual([f["invoices"] for f in out], [["INV-4"]])

    def test_vendor_spellings_grouped(self):
        rows = [{"invoice_number": str(i), "vendor_name": v, "sea_freight": 10} for i, v in enumerate(
            ["MEGALO FREIGHT", "MEGALO FREIGHT Cargo", "Megalo Freight Pvt. Ltd.", "Dachser India Pvt Ltd", "DSV Air & Sea"])]
        out = insights.vendor_aliases(frame(rows))
        self.assertEqual(len(out), 1)
        self.assertEqual(len(out[0]["metric"]["names"]), 3)

    def test_findings_note_counts_every_severity(self):
        fs = [{"severity": s} for s in ("high", "high", "medium", "low", "low", "low")]
        self.assertEqual(insights.severity_note(fs), "2 high, 1 medium, 3 low")
        self.assertEqual(insights.severity_note([{"severity": "low"}]), "1 low")
        self.assertEqual(insights.severity_note([]), "")


class Reports(unittest.TestCase):
    def test_workbook_and_dashboard_are_written(self):
        import openpyxl

        s = insights.analyse({"invoices": frame(many(60)), "runs": None, "source": {"host": "test"}},
                             now=pd.Timestamp("2026-12-15"))
        with tempfile.TemporaryDirectory() as tmp:
            x = report_xlsx.build(s, Path(tmp) / "mis.xlsx")
            wb = openpyxl.load_workbook(x)
            self.assertEqual(wb.sheetnames[:4], ["Summary", "Monthly", "Vendors", "Categories"])
            self.assertEqual(wb["Vendors"].max_row, 4)                      # header + 3 vendors
            self.assertTrue(len(wb["Monthly"]._charts) == 1)
            h = report_html.build(s, Path(tmp) / "d.html").read_text()
            self.assertIn('id="data"', h)
            self.assertNotIn("</script><script", h.split('id="data"')[1].split("</script>")[0])
            self.assertIn(s["meta"]["fingerprint"], h)


if __name__ == "__main__":
    unittest.main()
